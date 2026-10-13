from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

class ExcelGenerator:
    def __init__(self, file_name):
        self.file_name = file_name


     # 1. COLLECT ALL UNIQUE SUBJECTS
    def get_all_subjects(self,results):
        all_subjects = []
        for result in results:
            subjects = result["subjects"]
            for subject in subjects:
                subject_name = subject["subject_name"]
                if subject_name not in all_subjects:
                    all_subjects.append(subject_name)
        return all_subjects

    def calculate_sgpa(self, subjects):
        total_credit = 0
        total_credit_point = 0

        for subject in subjects:
            total_credit += float(subject["credit"])
            total_credit_point += float(subject["credit_point"])

        if total_credit == 0:
            return 0  # Avoid division by zero
        sgpa = total_credit_point / total_credit

        return round(sgpa,2)

    def create_excel(self, results, semester):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Student Results"

        # GET ALL SUBJECTS
        all_subject = self.get_all_subjects(results)

        #CREATE HEADER ROW
        headers = ["Register Number", "Student Name", "Semester"] + all_subject + ["SGPA"]
        sheet.append(headers)

        # Add student results to the sheet
        for result in results:
            student = result["student"]
            subject = result["subjects"]

            row = [
                student["register_number"],
                student["student_name"],
                semester
            ]

            subject_grade_map = {sub["subject_name"]: sub["grade"] for sub in subject}
            # subject_grade_map = {}
            # for subject in subjects:

            #     subject_grade_map[
            #         subject["subject_name"]
            #     ] = subject["grade"]

            # ADD GRADES IN CORRECT SUBJECT COLUMN
            for subject_name in all_subject:
                grade = subject_grade_map.get(subject_name, "N/A")  # Default to "N/A" if subject not found
                row.append(grade)

            # Calculate SGPA
            sgpa = self.calculate_sgpa(subject)
            row.append(sgpa)
            sheet.append(row)

        # Format Header
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center",vertical = "center")

        #allign data
        for row in sheet.iter_rows(min_row = 2):
            for cell in row:
                cell.alignment = Alignment(horizontal = "center", vertical = "center")

        # Auto column width
        for column in sheet.columns:
            maximum_length = 0
            column_letter = (column[0].column_letter)
            for cell in column:
                if cell.value is not None:
                    cell_length = len(str(cell.value))
                    if cell_length > maximum_length:
                        maximum_length = (cell_length)
            sheet.column_dimensions[column_letter].width = maximum_length + 2

        #Save excel
        workbook.save(self.file_name)
        print(f"\n Excel file: {self.file_name} is created sucessfully")


    # def create_excel(self, student, subjects, semester):

    #     # create a new workbook and select the active worksheet
    #     workbook = Workbook()

    #     sheet = workbook.active
    #     sheet.title = "Student Results"

    #     # Header row
    #     headers = [
    #         "Register Number",
    #         "Student Name",
    #         "Semster"
    #     ]

    #     for subject in subjects:
    #         headers.append(subject["subject_name"])

    #     headers.append("SGPA")

    #     # add the headers to the first row of the sheet
    #     sheet.append(headers)

    #     # STUDENT INFO ROW
    #     row =[
    #         student["register_number"],
    #         student["student_name"],
    #         semester
    #     ]

    #     for subject in subjects:
    #         row.append(subject["grade"])

    #     # Calculate SGPA
    #     total_credit = 0
    #     total_credit_point = 0

    #     for subject in subjects:
    #         total_credit += float(subject["credit"])
    #         total_credit_point += float(subject["credit_point"])

    #     if total_credit > 0 :
    #         sgpa = (total_credit_point/total_credit)

    #     else:
    #         sgpa = 0

    #     row.append(sgpa)

    #     sheet.append(row)

    #     workbook.save(self.file_name)

    #     print(f"Excel file '{self.file_name}' created successfully.")



        